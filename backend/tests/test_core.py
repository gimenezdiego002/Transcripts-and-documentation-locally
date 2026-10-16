from pathlib import Path
from time import monotonic, sleep

from app.jobs import JobManager
from app.extractors import extract_document
from app.organizer import generate_master
from app.storage import WorkspaceStore, atomic_write_text


def test_duplicate_detection_and_master(tmp_path: Path):
    store = WorkspaceStore(tmp_path / "knowledge")
    first_file = tmp_path / "notes.txt"
    first_file.write_text("Decision: ship the local version on 2026-10-16.\nOpen question: model size.", encoding="utf-8")
    with first_file.open("rb") as handle:
        source, duplicate = store.import_stream("notes.txt", handle)
    assert source is not None and duplicate is None
    with first_file.open("rb") as handle:
        second_source, duplicate = store.import_stream("copy.txt", handle)
    assert second_source is None
    assert duplicate and duplicate.source_id == source.source_id

    output = store.root / "extracted" / f"{source.source_id}.md"
    atomic_write_text(output, first_file.read_text(encoding="utf-8"))
    source.content_path = output.relative_to(store.root).as_posix()
    source.status = "completed"
    store.save_source(source)
    path, report = generate_master(store)
    master = path.read_text(encoding="utf-8")
    assert "Complete Original Content" in master
    assert source.source_id in master
    assert "2026-10-16" in master
    assert report["source_count"] == 1


def test_rejects_unsupported_file(tmp_path: Path):
    store = WorkspaceStore(tmp_path / "knowledge")
    unsupported = tmp_path / "payload.exe"
    unsupported.write_bytes(b"nope")
    with unsupported.open("rb") as handle:
        try:
            store.import_stream("payload.exe", handle)
        except ValueError as exc:
            assert "Unsupported" in str(exc)
        else:
            raise AssertionError("Unsupported extension was accepted")


def test_text_pipeline_job_generates_artifacts(tmp_path: Path):
    store = WorkspaceStore(tmp_path / "knowledge")
    source_file = tmp_path / "meeting.txt"
    source_file.write_text("Project Atlas\n\nWe decided to keep all files local.", encoding="utf-8")
    with source_file.open("rb") as handle:
        source, _ = store.import_stream(source_file.name, handle)
    assert source is not None

    manager = JobManager(store, lambda: "balanced")
    job = manager.submit(source.source_id)
    deadline = monotonic() + 5
    while manager.jobs[job.job_id].status not in {"completed", "failed"} and monotonic() < deadline:
        sleep(0.02)
    result = manager.jobs[job.job_id]
    manager.executor.shutdown(wait=True)

    assert result.status == "completed", result.error
    processed = store.get_source(source.source_id)
    assert processed.status == "completed"
    assert processed.content_path
    assert "Project Atlas" in store.read_content(processed)
    assert (store.root / "exports" / "MASTER_KNOWLEDGE.md").exists()


def test_archived_ids_are_not_reused(tmp_path: Path):
    store = WorkspaceStore(tmp_path / "knowledge")
    first = tmp_path / "first.txt"
    second = tmp_path / "second.txt"
    first.write_text("first", encoding="utf-8")
    second.write_text("second", encoding="utf-8")
    with first.open("rb") as handle:
        source, _ = store.import_stream(first.name, handle)
    assert source is not None
    store.archive_source(source.source_id)
    with second.open("rb") as handle:
        next_source, _ = store.import_stream(second.name, handle)
    assert next_source is not None
    assert next_source.source_id == "SRC_0002"


def test_docx_and_pdf_extraction(tmp_path: Path):
    import fitz
    from docx import Document

    docx_path = tmp_path / "requirements.docx"
    document = Document()
    document.add_heading("Project Requirements", level=1)
    document.add_paragraph("Keep all source files local.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Diego"
    table.cell(1, 1).text = "Approved"
    document.save(docx_path)

    docx_result = extract_document(docx_path)
    assert "# Project Requirements" in docx_result.markdown
    assert "Keep all source files local." in docx_result.markdown
    assert "| Owner | Status |" in docx_result.markdown

    pdf_path = tmp_path / "meeting.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Meeting decision: retain original recordings.")
    pdf.save(pdf_path)
    pdf.close()

    pdf_result = extract_document(pdf_path)
    assert "## Page 1" in pdf_result.markdown
    assert "retain original recordings" in pdf_result.markdown

